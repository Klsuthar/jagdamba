import openpyxl
import json
import re

excel_path = r"C:\Users\klsut\Documents\STUDENT.xlsx"
wb = openpyxl.load_workbook(excel_path)

def norm(s):
    if not s:
        return ""
    s = str(s).strip().upper()
    return re.sub(r'[^A-Z0-9]', '', s)

# Load English Medium
ws_eng = wb["English Medium"]
eng_students = []
for r in range(2, ws_eng.max_row + 1):
    sno = ws_eng.cell(r, 1).value
    cls = ws_eng.cell(r, 2).value
    sname = ws_eng.cell(r, 3).value
    fname = ws_eng.cell(r, 4).value
    mname = ws_eng.cell(r, 5).value
    dob = ws_eng.cell(r, 6).value
    dob_str = dob.strftime('%d/%m/%Y') if hasattr(dob, 'strftime') else (str(dob) if dob else '')
    eng_students.append({
        'sno': sno,
        'class': cls,
        'student_name': sname,
        'father_name': fname,
        'mother_name': mname,
        'dob': dob_str,
        'norm': norm(sname),
        'medium': 'English'
    })

# Load Hindi Medium
ws_hin = wb["Hindi Medium"]
hin_students = []
for r in range(2, ws_hin.max_row + 1):
    sno = ws_hin.cell(r, 1).value
    cls = ws_hin.cell(r, 2).value
    sname = ws_hin.cell(r, 3).value
    fname = ws_hin.cell(r, 4).value
    mname = ws_hin.cell(r, 5).value
    dob = ws_hin.cell(r, 6).value
    dob_str = dob.strftime('%d/%m/%Y') if hasattr(dob, 'strftime') else (str(dob) if dob else '')
    hin_students.append({
        'sno': sno,
        'class': cls,
        'student_name': sname,
        'father_name': fname,
        'mother_name': mname,
        'dob': dob_str,
        'norm': norm(sname),
        'medium': 'Hindi'
    })

class_map = {1: '1st', 2: '2nd', 3: '3rd', 4: '4th', 5: '5th'}

detailed_report = {}

for c in range(1, 6):
    json_path = rf"c:\Users\klsut\Documents\jagdamba\public\json\class{c}\class{c}_students.json"
    with open(json_path, encoding='utf-8') as f:
        site_students = json.load(f)
    
    cls_label = class_map[c]
    cls_results = []
    
    for stu in site_students:
        s_roll = stu['roll_no']
        s_name = stu['student_name']
        s_norm = norm(s_name)
        
        # 1. Matches in English Medium (same class)
        eng_same = [e for e in eng_students if e['class'] == cls_label and e['norm'] == s_norm]
        
        # 2. Matches in English Medium (other classes)
        eng_other = [e for e in eng_students if e['class'] != cls_label and e['norm'] == s_norm]
        
        # 3. Partial matches in English Medium (same class)
        eng_part_same = [e for e in eng_students if e['class'] == cls_label and e['norm'] != s_norm and (s_norm in e['norm'] or e['norm'] in s_norm)]
        
        # 4. Partial matches in English Medium (other classes)
        eng_part_other = [e for e in eng_students if e['class'] != cls_label and e['norm'] != s_norm and (s_norm in e['norm'] or e['norm'] in s_norm)]
        
        # 5. Matches in Hindi Medium (same class)
        hin_same = [h for h in hin_students if h['class'] == cls_label and (h['norm'] == s_norm or s_norm in h['norm'] or h['norm'] in s_norm)]
        
        # 6. Matches in Hindi Medium (other classes)
        hin_other = [h for h in hin_students if h['class'] != cls_label and (h['norm'] == s_norm or s_norm in h['norm'] or h['norm'] in s_norm)]
        
        cls_results.append({
            'roll': s_roll,
            'name': s_name,
            'eng_same': eng_same,
            'eng_other': eng_other,
            'eng_part_same': eng_part_same,
            'eng_part_other': eng_part_other,
            'hin_same': hin_same,
            'hin_other': hin_other,
        })
        
    detailed_report[c] = cls_results

# Print structured breakdown
for c in range(1, 6):
    print(f"\n=======================================================")
    print(f"CLASS {c} (Total: {len(detailed_report[c])} students)")
    print(f"=======================================================")
    
    clean = []
    need_review = []
    
    for r in detailed_report[c]:
        if len(r['eng_same']) == 1:
            clean.append(r)
        else:
            need_review.append(r)
            
    print(f"Clean single-match in same class (English Medium): {len(clean)}")
    print(f"Need Review / Clarification: {len(need_review)}")
    
    for r in need_review:
        print(f"\n  [Roll {r['roll']}] '{r['name']}':")
        if len(r['eng_same']) > 1:
            print(f"    - MULTIPLE SAME CLASS ENGLISH MATCHES:")
            for m in r['eng_same']:
                print(f"        * S.No {m['sno']}: {m['student_name']} | Father: {m['father_name']} | Mother: {m['mother_name']} | DOB: {m['dob']}")
        if r['eng_other']:
            print(f"    - ENGLISH MATCH IN OTHER CLASS:")
            for m in r['eng_other']:
                print(f"        * Class {m['class']}, S.No {m['sno']}: {m['student_name']} | Father: {m['father_name']} | Mother: {m['mother_name']} | DOB: {m['dob']}")
        if r['eng_part_same']:
            print(f"    - PARTIAL ENGLISH MATCH (SAME CLASS):")
            for m in r['eng_part_same']:
                print(f"        * S.No {m['sno']}: {m['student_name']} | Father: {m['father_name']} | Mother: {m['mother_name']} | DOB: {m['dob']}")
        if r['eng_part_other']:
            print(f"    - PARTIAL ENGLISH MATCH (OTHER CLASS):")
            for m in r['eng_part_other']:
                print(f"        * Class {m['class']}, S.No {m['sno']}: {m['student_name']} | Father: {m['father_name']} | Mother: {m['mother_name']} | DOB: {m['dob']}")
        if r['hin_same']:
            print(f"    - HINDI MEDIUM MATCH (SAME CLASS):")
            for m in r['hin_same']:
                print(f"        * S.No {m['sno']}: {m['student_name']} | Father: {m['father_name']} | Mother: {m['mother_name']} | DOB: {m['dob']}")
        if r['hin_other']:
            print(f"    - HINDI MEDIUM MATCH (OTHER CLASS):")
            for m in r['hin_other']:
                print(f"        * Class {m['class']}, S.No {m['sno']}: {m['student_name']} | Father: {m['father_name']} | Mother: {m['mother_name']} | DOB: {m['dob']}")
        if not r['eng_same'] and not r['eng_other'] and not r['eng_part_same'] and not r['eng_part_other'] and not r['hin_same'] and not r['hin_other']:
            print(f"    - *** NOT FOUND IN EXCEL (NEITHER ENGLISH NOR HINDI) ***")
