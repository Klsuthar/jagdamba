import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\klsut\Documents\STUDENT.xlsx")

queries = ["RAJ", "VARDHAN", "SHARMA", "TEENA", "TINA", "BHAM", "NEHRA", "NEHARA", "PRAJAPAT", "SWAMI"]

all_rows = []
for sname in ["English Medium", "Hindi Medium"]:
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        all_rows.append({
            'medium': sname,
            'sno': ws.cell(r, 1).value,
            'class': ws.cell(r, 2).value,
            'name': str(ws.cell(r, 3).value),
            'father': str(ws.cell(r, 4).value),
            'mother': str(ws.cell(r, 5).value),
            'dob': ws.cell(r, 6).value
        })

for q in queries:
    matches = [r for r in all_rows if q.upper() in r['name'].upper() or q.upper() in r['father'].upper()]
    print(f"\n--- Query '{q}' ({len(matches)} matches) ---")
    for m in matches:
        print(f"  [{m['medium']} | {m['class']}] {m['name']} (Father: {m['father']}, Mother: {m['mother']}, DOB: {m['dob']})")
