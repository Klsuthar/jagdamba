import json
import openpyxl
import os

# Updates to apply:
# Class 1, Roll 15: VINOD KHERIYA -> Father: DULARAM KHERIYA
# Class 2, Roll 2:  DEEPIKA -> Father: TARACHAND MEGHWAL, Mother: AMARI DEVI, DOB: 13/11/2019
# Class 2, Roll 15: TEENA SWAMI -> Father: GIRDHARI SWAMI

# 1. Update JSON files in public and dist
for base in ["public", "dist"]:
    # Class 1
    c1_path = rf"c:\Users\klsut\Documents\jagdamba\{base}\json\class1\class1_students.json"
    if os.path.exists(c1_path):
        with open(c1_path, "r", encoding="utf-8") as f:
            c1 = json.load(f)
        for s in c1:
            if s['roll_no'] == 15:
                s['father_name'] = "DULARAM KHERIYA"
        with open(c1_path, "w", encoding="utf-8") as f:
            json.dump(c1, f, indent=2, ensure_ascii=False)
        print(f"Updated Vinod Kheriya in {c1_path}")

    # Class 2
    c2_path = rf"c:\Users\klsut\Documents\jagdamba\{base}\json\class2\class2_students.json"
    if os.path.exists(c2_path):
        with open(c2_path, "r", encoding="utf-8") as f:
            c2 = json.load(f)
        for s in c2:
            if s['roll_no'] == 2:
                s['father_name'] = "TARACHAND MEGHWAL"
                s['mother_name'] = "AMARI DEVI"
                s['dob'] = "13/11/2019"
            elif s['roll_no'] == 15:
                s['father_name'] = "GIRDHARI SWAMI"
        with open(c2_path, "w", encoding="utf-8") as f:
            json.dump(c2, f, indent=2, ensure_ascii=False)
        print(f"Updated Deepika & Teena Swami in {c2_path}")

# 2. Update Student_List_All_Classes.xlsx in root, public, dist
excel_files = [
    r"c:\Users\klsut\Documents\jagdamba\Student_List_All_Classes.xlsx",
    r"c:\Users\klsut\Documents\jagdamba\public\data\Student_List_All_Classes.xlsx",
    r"c:\Users\klsut\Documents\jagdamba\dist\data\Student_List_All_Classes.xlsx"
]

for expath in excel_files:
    if not os.path.exists(expath):
        continue
    wb = openpyxl.load_workbook(expath)
    
    # Class 1
    ws1 = wb["Class 1"]
    h1 = [ws1.cell(1, col).value for col in range(1, ws1.max_column + 1)]
    idx_roll1 = h1.index('roll_no') + 1
    idx_f1 = h1.index('father_name') + 1
    for r in range(2, ws1.max_row + 1):
        if ws1.cell(r, idx_roll1).value == 15:
            ws1.cell(r, idx_f1, value="DULARAM KHERIYA")
            
    # Class 2
    ws2 = wb["Class 2"]
    h2 = [ws2.cell(1, col).value for col in range(1, ws2.max_column + 1)]
    idx_roll2 = h2.index('roll_no') + 1
    idx_f2 = h2.index('father_name') + 1
    idx_m2 = h2.index('mother_name') + 1
    idx_d2 = h2.index('dob') + 1
    for r in range(2, ws2.max_row + 1):
        roll_val = ws2.cell(r, idx_roll2).value
        if roll_val == 2:
            ws2.cell(r, idx_f2, value="TARACHAND MEGHWAL")
            ws2.cell(r, idx_m2, value="AMARI DEVI")
            ws2.cell(r, idx_d2, value="13/11/2019")
        elif roll_val == 15:
            ws2.cell(r, idx_f2, value="GIRDHARI SWAMI")
            
    wb.save(expath)
    print(f"Updated Excel file: {expath}")

print("\nAll 3 updates completed successfully!")
