import openpyxl
import json
import re

excel_path = r"C:\Users\klsut\Documents\STUDENT.xlsx"
wb = openpyxl.load_workbook(excel_path)

def norm(s):
    if not s:
        return ""
    s = str(s).strip().upper()
    return re.sub(r'[^A-Z0-9]', '', s)

# Load English Medium
ws_eng = wb["English Medium"]
eng_students = []
for r in range(2, ws_eng.max_row + 1):
    sno = ws_eng.cell(r, 1).value
    cls = ws_eng.cell(r, 2).value
    sname = ws_eng.cell(r, 3).value
    fname = ws_eng.cell(r, 4).value
    mname = ws_eng.cell(r, 5).value
    dob = ws_eng.cell(r, 6).value
    dob_str = dob.strftime('%d/%m/%Y') if hasattr(dob, 'strftime') else (str(dob) if dob else '')
    eng_students.append({
        'sno': sno,
        'class': cls,
        'student_name': sname,
        'father_name': fname,
        'mother_name': mname,
        'dob': dob_str,
        'norm': norm(sname),
        'medium': 'English'
    })

# Load Hindi Medium
ws_hin = wb["Hindi Medium"]
hin_students = []
for r in range(2, ws_hin.max_row + 1):
    sno = ws_hin.cell(r, 1).value
    cls = ws_hin.cell(r, 2).value
    sname = ws_hin.cell(r, 3).value
    fname = ws_hin.cell(r, 4).value
    mname = ws_hin.cell(r, 5).value
    dob = ws_hin.cell(r, 6).value
    dob_str = dob.strftime('%d/%m/%Y') if hasattr(dob, 'strftime') else (str(dob) if dob else '')
    hin_students.append({
        'sno': sno,
        'class': cls,
        'student_name': sname,
        'father_name': fname,
        'mother_name': mname,
        'dob': dob_str,
        'norm': norm(sname),
        'medium': 'Hindi'
    })

all_excel_students = eng_students + hin_students

class_map = {1: '1st', 2: '2nd', 3: '3rd', 4: '4th', 5: '5th'}

full_audit = []

for c in range(1, 6):
    json_path = rf"c:\Users\klsut\Documents\jagdamba\public\json\class{c}\class{c}_students.json"
    with open(json_path, encoding='utf-8') as f:
        site_students = json.load(f)
    
    cls_label = class_map[c]
    
    for stu in site_students:
        roll = stu['roll_no']
        name = stu['student_name']
        n_name = norm(name)
        
        entry = {
            'class_num': c,
            'class_label': cls_label,
            'roll_no': roll,
            'site_name': name,
            'match_type': None,
            'excel_record': None,
            'candidates': [],
            'notes': ''
        }
        
        # 1. Exact match in English Medium of SAME class
        matches = [e for e in eng_students if e['class'] == cls_label and e['norm'] == n_name]
        if len(matches) == 1:
            entry['match_type'] = 'EXACT_ENGLISH_SAME_CLASS'
            entry['excel_record'] = matches[0]
            full_audit.append(entry)
            continue
        elif len(matches) > 1:
            entry['match_type'] = 'AMBIGUOUS_SAME_CLASS'
            entry['candidates'] = matches
            entry['notes'] = f"Found {len(matches)} students with exact same name in English Medium {cls_label}"
            full_audit.append(entry)
            continue
            
        # 2. Check known cross-class or spelling variations
        # Class 1:
        if c == 1:
            if n_name == norm('KARMVEER SINGH'):
                cand = [e for e in eng_students if e['class'] == '1st' and 'KARMVEER' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}'"
                    full_audit.append(entry); continue
            if n_name == norm('MAGHARAM BHADU'):
                cand = [e for e in eng_students if e['class'] == '1st' and 'MAGHARAM' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}' (BHAMU vs BHADU)"
                    full_audit.append(entry); continue
            if n_name == norm('RAJYAVARDHAN SINGH'):
                cand = [e for e in eng_students if e['class'] == '1st' and 'RAJAYVERDHAN' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}'"
                    full_audit.append(entry); continue
            if n_name == norm('PARVEEN KUMAR'):
                cand = [e for e in eng_students if e['class'] == '1st' and 'PRAVEENKUMAR' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}' (PRAVEEN vs PARVEEN)"
                    full_audit.append(entry); continue
            if n_name == norm('MINAXI SHARMA'):
                cand = [e for e in eng_students if e['class'] == '1st' and 'MIAKSHI' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}' (Father: RAM LAL SHARMA)"
                    full_audit.append(entry); continue
            if n_name == norm('GUNJAN'):
                cand = [e for e in eng_students if e['class'] == '1st' and 'GUNJAN' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}' (Father: HARIRAM SINWAL)"
                    full_audit.append(entry); continue
            if n_name == norm('SANTOSH KHAIRIYA'):
                cand = [e for e in eng_students if e['class'] == '1st' and e['norm'] == 'SANTOSH']
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}' (Father: NANDURAM)"
                    full_audit.append(entry); continue
            if n_name == norm('YOGENDRA'):
                cand = [e for e in eng_students if 'YOGENDRA' in e['norm']]
                entry['match_type'] = 'DOWNGRADED_CROSS_CLASS'
                entry['candidates'] = cand
                entry['notes'] = "Downgraded student? Found Yogendra Singh in 2nd and 3rd English Medium"
                full_audit.append(entry); continue
            if n_name == norm('NARESH'):
                cand = [h for h in all_excel_students if 'NARESH' in h['norm']]
                entry['match_type'] = 'HINDI_OR_CROSS_CLASS'
                entry['candidates'] = cand
                entry['notes'] = "Found Naresh Suthar (4th Hindi)"
                full_audit.append(entry); continue
            if n_name == norm('DHARMENDRA'):
                cand = [h for h in all_excel_students if 'DHARMENDRA' in h['norm']]
                entry['match_type'] = 'HINDI_OR_CROSS_CLASS'
                entry['candidates'] = cand
                entry['notes'] = "Found Dharmendra Simar (4th Hindi)"
                full_audit.append(entry); continue
            if n_name == norm('MAHAK SINWAL'):
                cand = [h for h in all_excel_students if 'MAHAK' in h['norm']]
                entry['match_type'] = 'HINDI_OR_CROSS_CLASS'
                entry['candidates'] = cand
                entry['notes'] = "Found Mahak (5th Hindi, Father: Shankar Nath)"
                full_audit.append(entry); continue
            if n_name == norm('PANKAJ KHAIRIYA'):
                cand = [h for h in all_excel_students if 'PANKAJ' in h['norm']]
                entry['match_type'] = 'HINDI_OR_CROSS_CLASS'
                entry['candidates'] = cand
                entry['notes'] = "Found Pankaj (5th Hindi, Father: Shera Ram)"
                full_audit.append(entry); continue
            if n_name == norm('RAMDEV SUTHAR'):
                cand = [h for h in all_excel_students if 'RAMDEV' in h['norm']]
                entry['match_type'] = 'HINDI_OR_CROSS_CLASS'
                entry['candidates'] = cand
                entry['notes'] = "Found Ramdev Suthar (3rd Hindi, Father: Rajuram)"
                full_audit.append(entry); continue
            if n_name == norm('NEHA'):
                cand = [h for h in all_excel_students if 'NEHA' in h['norm'] or 'NEHARA' in h['norm']]
                entry['match_type'] = 'AMBIGUOUS'
                entry['candidates'] = cand
                entry['notes'] = "Multiple NEHA candidates in Excel"
                full_audit.append(entry); continue
            if n_name == norm('VINOD KHERIYA'):
                entry['match_type'] = 'NOT_FOUND'
                entry['notes'] = "Not found in Excel"
                full_audit.append(entry); continue

        # Class 2:
        if c == 2:
            if n_name == norm('GAJENDRA SINGH'):
                cand = [e for e in eng_students if e['class'] == '1st' and e['norm'] == n_name]
                if cand:
                    entry['match_type'] = 'CROSS_CLASS_ENGLISH'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Student is in Class 1st in English Medium"
                    full_audit.append(entry); continue
            if n_name == norm('DHARMENDR PRAJAPAT'):
                cand = [h for h in hin_students if h['class'] == '2nd' and h['norm'] == n_name]
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 2nd"
                    full_audit.append(entry); continue
            if n_name == norm('DIVYA SUTHAR'):
                cand = [h for h in all_excel_students if 'DIVYA' in h['norm']]
                entry['match_type'] = 'AMBIGUOUS_HINDI_CROSS'
                entry['candidates'] = cand
                entry['notes'] = "In Hindi 2nd: Divya (F: Girdhari Lal); In Eng 4th: Divya (F: Labhuram Kheriya)"
                full_audit.append(entry); continue
            if n_name == norm('TEENA SWAMI'):
                entry['match_type'] = 'NOT_FOUND'
                entry['notes'] = "Not found in Excel"
                full_audit.append(entry); continue

        # Class 3:
        if c == 3:
            if n_name == norm('VIJAY KHERIYA'):
                cand = [e for e in eng_students if e['class'] == '3rd' and 'VIJAY' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}' (KHERIA vs KHERIYA)"
                    full_audit.append(entry); continue
            if n_name == norm('DIKSHIT CHARAN'):
                cand = [e for e in eng_students if e['class'] == '3rd' and 'DIKSHEET' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}'"
                    full_audit.append(entry); continue
            if n_name == norm('AASISH'):
                cand = [e for e in eng_students if e['class'] == '3rd' and 'AASHISH' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}'"
                    full_audit.append(entry); continue
            if n_name == norm('SUNITA KHERIYA'):
                cand = [h for h in hin_students if h['class'] == '3rd' and 'SUNITA' in h['norm']]
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 3rd"
                    full_audit.append(entry); continue
            if n_name == norm('BABULAL'):
                cand = [h for h in hin_students if h['class'] == '3rd' and 'BABU' in h['norm']]
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 3rd (Babu Lal, F: Girdhari Lal)"
                    full_audit.append(entry); continue
            if n_name == norm('NAITIK SONI'):
                cand = [h for h in hin_students if h['class'] == '3rd' and 'NAITIK' in h['norm']]
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 3rd"
                    full_audit.append(entry); continue
            if n_name == norm('PANKAJ NEHARA'):
                cand = [h for h in hin_students if h['class'] == '3rd' and 'PANKAJ' in h['norm']]
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 3rd (Pankaj Nehra, F: Sahiram Nehra)"
                    full_audit.append(entry); continue
            if n_name == norm('NEHA BHAMU'):
                cand = [h for h in hin_students if h['class'] == '3rd' and 'NEHA' in h['norm']]
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 3rd (Neha Bhambhu, F: Sanjay Kumar)"
                    full_audit.append(entry); continue

        # Class 4:
        if c == 4:
            if n_name == norm('NISHTHA'):
                cand = [e for e in eng_students if e['class'] == '4th' and 'NISHTHA' in e['norm']]
                if cand:
                    entry['match_type'] = 'VARIATION_ENGLISH_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = f"Matched to '{cand[0]['student_name']}'"
                    full_audit.append(entry); continue
            if n_name == norm('NARESH SUTHAR'):
                cand = [h for h in hin_students if h['class'] == '4th' and 'NARESH' in h['norm']]
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 4th (Naresh Suthar, F: Girdhari Lal)"
                    full_audit.append(entry); continue
            if n_name == norm('PRAVEENA'):
                cand = [h for h in hin_students if h['class'] == '4th' and 'PRAVEENA' in h['norm']]
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 4th (Praveena, F: Govind Ram Jat)"
                    full_audit.append(entry); continue

        # Class 5:
        if c == 5:
            # DINDAYAL, DINESH SIMAR, KAUSHALYA, MANISHA, NAVEEN SIMAR are in English Medium Class 4th!
            for orig_eng in [('DINDAYAL', 'DINDAYAL'), ('DINESH SIMAR', 'DINESH SIMAR'), ('KAUSHALYA', 'KAUSHALYA'), ('MANISHA', 'MANISHA'), ('NAVEEN SIMAR', 'NAVEEN SIMAR')]:
                if n_name == norm(orig_eng[0]):
                    cand = [e for e in eng_students if e['class'] == '4th' and e['norm'] == norm(orig_eng[1])]
                    if cand:
                        entry['match_type'] = 'CROSS_CLASS_ENGLISH_4TH'
                        entry['excel_record'] = cand[0]
                        entry['notes'] = "Promoted/shifted from Class 4th English Medium to Class 5th"
                        break
            if entry['match_type']:
                full_audit.append(entry); continue
                
            if n_name == norm('LAXMI'):
                cand = [h for h in hin_students if h['class'] == '5th' and h['norm'] == 'LAXMI']
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 5th (Laxmi, F: Rakesh Suthar)"
                    full_audit.append(entry); continue
            if n_name == norm('SONAKSHI KANWAR'):
                cand = [h for h in hin_students if h['class'] == '5th' and h['norm'] == norm('SONAKSHI KANWAR')]
                if cand:
                    entry['match_type'] = 'HINDI_MEDIUM_SAME_CLASS'
                    entry['excel_record'] = cand[0]
                    entry['notes'] = "Found in Hindi Medium Class 5th (Sonakshi Kanwar, F: Jitendra Singh)"
                    full_audit.append(entry); continue

        # Catch-all
        entry['match_type'] = 'UNRESOLVED'
        entry['notes'] = f"No match rules triggered for '{name}'"
        full_audit.append(entry)

print("Full audit completed!")
from collections import Counter
print(Counter(r['match_type'] for r in full_audit))
