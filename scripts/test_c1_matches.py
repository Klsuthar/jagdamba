import openpyxl
import json
import re

excel_path = r"C:\Users\klsut\Documents\STUDENT.xlsx"
wb = openpyxl.load_workbook(excel_path)

def norm(s):
    if not s:
        return ""
    s = str(s).strip().upper()
    return re.sub(r'[^A-Z0-9]', '', s)

# Load English Medium
ws_eng = wb["English Medium"]
eng_students = []
for r in range(2, ws_eng.max_row + 1):
    sno = ws_eng.cell(r, 1).value
    cls = ws_eng.cell(r, 2).value
    sname = ws_eng.cell(r, 3).value
    fname = ws_eng.cell(r, 4).value
    mname = ws_eng.cell(r, 5).value
    dob = ws_eng.cell(r, 6).value
    dob_str = dob.strftime('%d/%m/%Y') if hasattr(dob, 'strftime') else (str(dob) if dob else '')
    eng_students.append({
        'sno': sno,
        'class': cls,
        'student_name': sname,
        'father_name': fname,
        'mother_name': mname,
        'dob': dob_str,
        'norm': norm(sname),
        'medium': 'English'
    })

# Load Hindi Medium
ws_hin = wb["Hindi Medium"]
hin_students = []
for r in range(2, ws_hin.max_row + 1):
    sno = ws_hin.cell(r, 1).value
    cls = ws_hin.cell(r, 2).value
    sname = ws_hin.cell(r, 3).value
    fname = ws_hin.cell(r, 4).value
    mname = ws_hin.cell(r, 5).value
    dob = ws_hin.cell(r, 6).value
    dob_str = dob.strftime('%d/%m/%Y') if hasattr(dob, 'strftime') else (str(dob) if dob else '')
    hin_students.append({
        'sno': sno,
        'class': cls,
        'student_name': sname,
        'father_name': fname,
        'mother_name': mname,
        'dob': dob_str,
        'norm': norm(sname),
        'medium': 'Hindi'
    })

# Now let's examine each class carefully:
results = {}

for c in range(1, 6):
    json_path = rf"c:\Users\klsut\Documents\jagdamba\public\json\class{c}\class{c}_students.json"
    with open(json_path, encoding='utf-8') as f:
        site_students = json.load(f)
    results[c] = site_students

print("=== CLASS 1 (26 students) ===")
for s in results[1]:
    r = s['roll_no']
    name = s['student_name']
    n = norm(name)
    m = [e for e in eng_students if e['class'] == '1st' and e['norm'] == n]
    if m:
        print(f"Roll {r:2d}: {name:20s} -> MATCH: {m[0]['student_name']} (Father: {m[0]['father_name']}, Mother: {m[0]['mother_name']}, DOB: {m[0]['dob']})")
    else:
        print(f"Roll {r:2d}: {name:20s} -> NO EXACT MATCH IN 1st ENG")
